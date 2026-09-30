"""Читання CV з PDF або DOCX."""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader


class CVReadError(ValueError):
    pass


@dataclass
class CVFile:
    filename: str
    text: str  # Витягнутий текст: для застосування правок і для Grill me.
    pdf_bytes: bytes | None = None  # Оригінал PDF, щоб модель бачила й верстку.

    def as_content_blocks(self) -> list[dict]:
        """CV у форматі блоків для запиту до Claude."""
        if self.pdf_bytes is not None:
            return [
                {
                    "type": "document",
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": base64.standard_b64encode(self.pdf_bytes).decode("ascii"),
                    },
                    "title": self.filename,
                }
            ]
        return [{"type": "text", "text": f"<cv filename=\"{self.filename}\">\n{self.text}\n</cv>"}]


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as e:  # pypdf кидає різні типи помилок на битих файлах
        raise CVReadError("Не вдалося прочитати PDF. Можливо, файл пошкоджений.") from e


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise CVReadError("Не вдалося прочитати DOCX. Можливо, файл пошкоджений.") from e
    lines = [p.text for p in doc.paragraphs]
    # Багато шаблонів CV тримають текст у таблицях.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines).strip()


def load_cv(filename: str, data: bytes) -> CVFile:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _pdf_text(data)
        # Скановані PDF не мають тексту, але модель усе одно прочитає їх як картинку.
        return CVFile(filename=filename, text=text, pdf_bytes=data)
    if name.endswith(".docx"):
        text = _docx_text(data)
        if not text:
            raise CVReadError("У файлі DOCX немає тексту.")
        return CVFile(filename=filename, text=text)
    if name.endswith(".doc"):
        raise CVReadError("Старий формат .doc не підтримується. Збережи файл як .docx або .pdf.")
    raise CVReadError("Підтримуються тільки PDF і DOCX.")
